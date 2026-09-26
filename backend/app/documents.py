from __future__ import annotations

import hashlib
import io
import os
import re
from functools import lru_cache
from pathlib import Path

from docx import Document as DocxDocument
from pypdf import PdfReader

from .models import DocumentChunk, DocumentRole, PolicyDocument


PRELOADED_POLICIES = {
    "abhi-activ-one": {
        "filename": "ABHI Product Brochure.pdf",
        "insurer": "Aditya Birla Health Insurance",
        "product": "Activ One",
    },
    "care-health": {
        "filename": "Care Health Product Brochure.pdf",
        "insurer": "Care Health Insurance",
        "product": "Care Health",
    },
    "hdfc-optima-secure": {
        "filename": "HDFC Product Brochure.pdf",
        "insurer": "HDFC ERGO",
        "product": "Optima Secure",
    },
    "niva-reassure-2": {
        "filename": "Niva Bupa Product Brochure.pdf",
        "insurer": "Niva Bupa",
        "product": "ReAssure 2.0",
    },
}


def _clean(text: str) -> str:
    return re.sub(r"\s+", " ", text).strip()


def _split_text(text: str, size: int = 1200, overlap: int = 160) -> list[str]:
    text = _clean(text)
    if not text:
        return []
    chunks: list[str] = []
    start = 0
    while start < len(text):
        end = min(len(text), start + size)
        if end < len(text):
            boundary = text.rfind(". ", start, end)
            if boundary > start + size // 2:
                end = boundary + 1
        chunks.append(text[start:end].strip())
        if end == len(text):
            break
        start = max(start + 1, end - overlap)
    return chunks


def parse_pdf_bytes(
    data: bytes,
    document_id: str,
    name: str,
    role: DocumentRole,
    insurer: str | None = None,
    product_name: str | None = None,
) -> PolicyDocument:
    reader = PdfReader(io.BytesIO(data))
    if reader.is_encrypted and reader.decrypt("") == 0:
        raise ValueError(f"Password-protected PDF is not supported: {name}")
    chunks: list[DocumentChunk] = []
    for page_number, page in enumerate(reader.pages, start=1):
        for index, text in enumerate(_split_text(page.extract_text() or "")):
            chunks.append(
                DocumentChunk(
                    chunk_id=f"{document_id}:p{page_number}:c{index + 1}",
                    page_number=page_number,
                    text=text,
                )
            )
    if not chunks:
        raise ValueError(f"No extractable text found in {name}")
    return PolicyDocument(
        document_id=document_id,
        document_name=name,
        role=role,
        insurer=insurer,
        product_name=product_name,
        file_hash=hashlib.sha256(data).hexdigest(),
        chunks=chunks,
    )


def parse_docx_bytes(
    data: bytes,
    document_id: str,
    name: str,
    role: DocumentRole,
) -> PolicyDocument:
    document = DocxDocument(io.BytesIO(data))
    paragraphs = [_clean(paragraph.text) for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            paragraphs.append(" | ".join(_clean(cell.text) for cell in row.cells))
    chunks = [
        DocumentChunk(chunk_id=f"{document_id}:s{index + 1}", section=f"Section {index + 1}", text=text)
        for index, text in enumerate(_split_text("\n".join(filter(None, paragraphs))))
    ]
    if not chunks:
        raise ValueError(f"No extractable text found in {name}")
    return PolicyDocument(
        document_id=document_id,
        document_name=name,
        role=role,
        file_hash=hashlib.sha256(data).hexdigest(),
        chunks=chunks,
    )


def load_preloaded_policies(policy_directory: Path) -> list[PolicyDocument]:
    documents: list[PolicyDocument] = []
    for document_id, metadata in PRELOADED_POLICIES.items():
        path = policy_directory / metadata["filename"]
        if not path.exists():
            continue
        documents.append(
            parse_pdf_bytes(
                path.read_bytes(),
                document_id=document_id,
                name=metadata["filename"],
                role=DocumentRole.PRELOADED_POLICY,
                insurer=metadata["insurer"],
                product_name=metadata["product"],
            )
        )
    return documents


def default_policy_directory() -> Path:
    configured = os.getenv("POLICY_DIRECTORY")
    if configured:
        return Path(configured).expanduser().resolve()
    bundled = Path(__file__).resolve().parents[1] / "data" / "policies"
    if bundled.exists():
        return bundled
    return Path(__file__).resolve().parents[2] / "Policy Documents"


@lru_cache(maxsize=1)
def cached_preloaded_policies() -> tuple[PolicyDocument, ...]:
    documents = load_preloaded_policies(default_policy_directory())
    if len(documents) != len(PRELOADED_POLICIES):
        loaded = {document.document_id for document in documents}
        missing = sorted(set(PRELOADED_POLICIES) - loaded)
        raise RuntimeError(f"Missing preloaded policy documents: {', '.join(missing)}")
    return tuple(documents)
