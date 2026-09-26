from __future__ import annotations

from io import BytesIO
import json
import os
from pathlib import Path
import shutil
import subprocess
import tempfile
from pptx import Presentation
from pptx.chart.data import ChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt
from .models import MarketingPitch

NAVY = RGBColor(15, 44, 73); BLUE = RGBColor(0, 112, 173); TEAL = RGBColor(0, 140, 149)
MINT = RGBColor(78, 229, 209); LIGHT = RGBColor(240, 245, 248); DARK = RGBColor(35, 43, 51); WHITE = RGBColor(255, 255, 255)

def _fit_text(value: object, maximum: int) -> str:
    text = " ".join(str(value).replace("...", ".").replace("…", ".").split())
    if len(text) <= maximum:
        return text
    complete = next((sentence.strip() for sentence in re.findall(r"[^.!?]+[.!?]", text) if len(sentence.strip()) <= maximum), "")
    if complete:
        return complete
    candidate = text[:maximum]
    boundary = candidate.rfind(" ")
    if boundary >= maximum * .55:
        candidate = candidate[:boundary]
    return candidate.rstrip(" ,:;-.") + "."

def _sanitise_pitch(pitch: MarketingPitch) -> MarketingPitch:
    def clean(value):
        if isinstance(value, str):
            return value.replace("...", ".").replace("…", ".")
        if isinstance(value, list):
            return [clean(item) for item in value]
        if isinstance(value, dict):
            return {key: clean(item) for key, item in value.items()}
        return value
    return MarketingPitch.model_validate(clean(pitch.model_dump(mode="python")))

def _validate_no_truncation_markers(payload: bytes) -> None:
    presentation = Presentation(BytesIO(payload))
    offenders = [shape.text for slide in presentation.slides for shape in slide.shapes if hasattr(shape, "text") and ("..." in shape.text or "…" in shape.text)]
    if offenders:
        raise ValueError("PPT export contains truncated text markers: " + " | ".join(offenders[:3]))

def _text(slide, text, x, y, w, h, size=17, color=DARK, bold=False):
    text = str(text).replace("\x07", "•")
    frame = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)).text_frame
    frame.word_wrap = True; frame.clear()
    for index, line in enumerate(str(text).split("\n")):
        p = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        p.text = line; p.font.name = "Aptos"; p.font.size = Pt(size); p.font.bold = bold; p.font.color.rgb = color

def _base(slide, title, company):
    slide.background.fill.solid(); slide.background.fill.fore_color.rgb = WHITE
    band = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(.16), Inches(7.5))
    band.fill.solid(); band.fill.fore_color.rgb = TEAL; band.line.fill.background()
    _text(slide, title, .55, .24, 9.7, .55, 28, NAVY, True)
    box = slide.shapes.add_textbox(Inches(10.35), Inches(.3), Inches(2.35), Inches(.35)); p = box.text_frame.paragraphs[0]
    p.text = company; p.alignment = PP_ALIGN.RIGHT; p.font.name = "Aptos"; p.font.size = Pt(10); p.font.color.rgb = BLUE

def _footer(slide, number, citations):
    rule = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(.55), Inches(7.01), Inches(12.2), Inches(.012))
    rule.fill.solid(); rule.fill.fore_color.rgb = RGBColor(214, 226, 232); rule.line.fill.background()
    label = "Sources: " + "; ".join(f"{s.document_name} - {('p.' + str(s.page_number)) if s.page_number else (s.section or 'location not established')}" for s in citations[:4]) if citations else "Sources: company profile, advisor inputs, and supplied policy documents"
    _text(slide, label, .58, 7.06, 11.35, .23, 8); _text(slide, f"{number:02d}", 12.15, 7.04, .5, .24, 9, TEAL, True)

def _bullets(slide, items, x, y, w, h, size=14):
    _text(slide, "\n".join(f"• {item}" for item in items[:8] if item), x, y, w, h, size)

def _card(slide, title, value, x, y):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(2.78), Inches(1.08))
    shape.fill.solid(); shape.fill.fore_color.rgb = LIGHT; shape.line.color.rgb = RGBColor(211, 225, 232)
    _text(slide, title.upper(), x+.2, y+.14, 2.4, .22, 9, TEAL, True); _text(slide, value, x+.2, y+.43, 2.4, .5, 15, NAVY, True)

def _chart(slide, chart, x=6.3, y=1.55, w=6.25, h=4.85):
    data = ChartData(); data.categories = chart.categories
    for name, values in chart.series.items(): data.add_series(name, values)
    native = slide.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(x), Inches(y), Inches(w), Inches(h), data).chart
    native.has_title = True; native.chart_title.text_frame.text = chart.title; native.has_legend = len(chart.series) > 1
    if native.has_legend: native.legend.position = XL_LEGEND_POSITION.BOTTOM
    native.value_axis.minimum_scale = 0; native.category_axis.tick_labels.font.size = Pt(10); native.value_axis.tick_labels.font.size = Pt(9)

def _table(slide, table):
    native = slide.shapes.add_table(len(table.rows)+1, len(table.columns), Inches(.55), Inches(1.4), Inches(12.15), Inches(5.35)).table
    for col, label in enumerate(table.columns):
        cell = native.cell(0, col); cell.text = label; cell.fill.solid(); cell.fill.fore_color.rgb = NAVY
        for p in cell.text_frame.paragraphs: p.font.size = Pt(10); p.font.bold = True; p.font.color.rgb = WHITE
    for row_index, row in enumerate(table.rows, 1):
        for col, value in enumerate(row):
            cell = native.cell(row_index, col); cell.text = _fit_text(value, 210); cell.fill.solid(); cell.fill.fore_color.rgb = LIGHT if row_index % 2 else WHITE
            for p in cell.text_frame.paragraphs: p.font.size = Pt(9); p.font.color.rgb = DARK

def _export_pitch_pptx_python(pitch: MarketingPitch) -> bytes:
    if len(pitch.pitch_slides) not in {3, 4, 5}: raise ValueError("Client presentation must contain three to five slides")
    deck = Presentation(); deck.slide_width = Inches(13.333); deck.slide_height = Inches(7.5); blank = deck.slide_layouts[6]
    for item in pitch.pitch_slides:
        slide = deck.slides.add_slide(blank); _base(slide, item.title, pitch.company_profile.company_name)
        _text(slide, item.key_message, .58, .92, 12, .48, 13); blocks = [value for block in item.content_blocks for value in block.items]
        if item.slide_number == 1:
            recommended = next((v.split(":",1)[1].strip() for v in blocks if v.startswith("Recommended policy:")), "Recommendation pending")
            callout = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(.6), Inches(1.58), Inches(5.25), Inches(1.22)); callout.fill.solid(); callout.fill.fore_color.rgb = NAVY; callout.line.fill.background()
            _text(slide, "RECOMMENDED POLICY", .9, 1.78, 4.6, .25, 10, MINT, True); _text(slide, recommended, .9, 2.05, 4.55, .5, 25, WHITE, True)
            _bullets(slide, [v for v in blocks if not v.startswith("Recommended policy:")], .72, 3.08, 5.15, 3.45)
            if item.charts: _chart(slide, item.charts[0], 6.25, 1.55, 6.35, 4.9)
        elif item.slide_number == 2:
            facts = item.content_blocks[0].items if item.content_blocks else []
            for i, fact in enumerate(facts[:4]):
                label, _, value = fact.partition(":"); _card(slide, label, value.strip(), .65+(i%2)*3, 1.55+(i//2)*1.35)
            _text(slide, "EXPOSURES AND REQUIREMENTS", 6.65, 1.58, 5.5, .3, 10, TEAL, True)
            _bullets(slide, [v for b in item.content_blocks[1:] for v in b.items], 6.65, 1.98, 5.65, 4.6)
        elif item.comparison_tables: _table(slide, item.comparison_tables[0])
        else:
            _bullets(slide, blocks, .72, 1.55, 5.35, 4.95)
            if item.charts: _chart(slide, item.charts[0])
        if item.slide_number == 5: _text(slide, "Advisor review and a PASS audit remain mandatory before client release.", .72, 6.56, 11.5, .24, 9)
        _footer(slide, item.slide_number, item.visible_citations)
    output = BytesIO(); deck.save(output); return output.getvalue()


def export_pitch_pptx(pitch: MarketingPitch) -> bytes:
    """Render through PptxGenJS; retain the tested Python renderer as a safe fallback."""
    pitch = _sanitise_pitch(pitch)
    engine = Path(__file__).resolve().parent.parent / "pptx_engine" / "render.mjs"
    node = os.getenv("NODE_BINARY") or shutil.which("node")
    if node and engine.exists():
        with tempfile.TemporaryDirectory(prefix="marsh-pptx-") as temp_dir:
            payload_path = Path(temp_dir) / "pitch.json"
            output_path = Path(temp_dir) / "pitch.pptx"
            payload_path.write_text(json.dumps(pitch.model_dump(mode="json"), ensure_ascii=False), encoding="utf-8")
            result = subprocess.run(
                [node, str(engine), str(payload_path), str(output_path)],
                capture_output=True, text=True, timeout=45, check=False,
            )
            if result.returncode == 0 and output_path.exists():
                payload = output_path.read_bytes()
                if payload[:2] == b"PK":
                    _validate_no_truncation_markers(payload)
                    return payload
    payload = _export_pitch_pptx_python(pitch)
    _validate_no_truncation_markers(payload)
    return payload
